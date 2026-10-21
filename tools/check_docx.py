from docx import Document

d = Document("整理.docx")
print("paragraphs", len(d.paragraphs), "tables", len(d.tables), "sections", len(d.sections))
print("--- headings ---")
for p in d.paragraphs:
    if p.style.name.startswith("Heading"):
        pad = "  " if p.style.name == "Heading 2" else ""
        print(pad + p.text)

tick = chr(96)
bad = [p.text for p in d.paragraphs if tick in p.text]
print("bad literal:", bad)

print("--- tables ---")
for i, t in enumerate(d.tables):
    head = " / ".join(c.text for c in t.rows[0].cells)
    print(i, len(t.rows), "x", len(t.columns), "|", head)
