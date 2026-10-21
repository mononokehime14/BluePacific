"""生成《整理.docx》：端侧 LLM 知识汇总的容器文档。

设计目标
- 这是一本"活页夹"：每个二级标题对应一份 md 笔记，用户把 md 内容粘到标题下面。
- 自带索引表（可当进度表用）、页码、统一的标题/正文/表格样式。
"""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_CN = "等线"
BODY_EN = "DengXian"
HEAD_CN = "微软雅黑"
CODE_EN = "Consolas"

C_DARK = RGBColor(0x1F, 0x38, 0x64)
C_MID = RGBColor(0x2E, 0x54, 0x96)
C_TEXT = RGBColor(0x26, 0x26, 0x26)
C_GRAY = RGBColor(0x80, 0x80, 0x80)
C_HINT = RGBColor(0x9A, 0x9A, 0x9A)


# ---------------------------------------------------------------- 基础工具
def set_run_font(run, size=None, bold=None, color=None, cn=BODY_CN, en=BODY_EN, italic=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)


def style_font(style, size=None, bold=None, color=None, cn=BODY_CN, en=BODY_EN):
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    style.font.name = en
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)


def bottom_border(paragraph, color="1F3864", size=8, space=4):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def cell_text(cell, text, size=9.5, bold=False, color=C_TEXT, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def _tbl_cell_margins(table, left=85, right=85):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("left", left), ("right", right), ("top", 0), ("bottom", 0)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    try:
        tblPr.insert_element_before(mar, "w:tblLook", "w:tblCaption", "w:tblDescription")
    except Exception:
        tblPr.append(mar)


def make_table(doc, header, rows, widths, header_fill="1F3864", body_size=9.5):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False  # -> <w:tblLayout w:type="fixed"/>
    _tbl_cell_margins(table)
    for i, w in enumerate(widths):
        table.columns[i].width = Cm(w)
    hdr = table.rows[0]
    for i, text in enumerate(header):
        cell = hdr.cells[i]
        cell_text(cell, text, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, header_fill)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            bold = i == 0
            cell_text(cells[i], str(text), size=body_size, bold=bold)
    # 固定列宽（tcW）+ 禁止行跨页断开
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
            tcpr = row.cells[i]._tc.get_or_add_tcPr()
            va = OxmlElement("w:vAlign")
            va.set(qn("w:val"), "center")
            tcpr.append(va)
    # 表头重复
    trpr = table.rows[0]._tr.get_or_add_trPr()
    tblheader = OxmlElement("w:tblHeader")
    tblheader.set(qn("w:val"), "true")
    trpr.append(tblheader)
    return table


def hint(doc, text):
    """浅灰小字提示：指出该标题对应的 md 文件，粘贴后可删。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=8.5, color=C_HINT, italic=True)


def part(doc, title):
    doc.add_page_break()
    h = doc.add_heading(title, level=1)
    return h


def section(doc, title):
    return doc.add_heading(title, level=2)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=9, color=C_GRAY)


# ---------------------------------------------------------------- 文档
doc = Document()

sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)

normal = doc.styles["Normal"]
style_font(normal, size=10.5, color=C_TEXT, cn=BODY_CN)
normal.paragraph_format.line_spacing = 1.42
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.space_before = Pt(0)

h1 = doc.styles["Heading 1"]
style_font(h1, size=18, bold=True, color=C_DARK, cn=HEAD_CN, en=HEAD_CN)
h1.paragraph_format.space_before = Pt(6)
h1.paragraph_format.space_after = Pt(14)
h1.paragraph_format.keep_with_next = True

h2 = doc.styles["Heading 2"]
style_font(h2, size=13.5, bold=True, color=C_MID, cn=HEAD_CN, en=HEAD_CN)
h2.paragraph_format.space_before = Pt(16)
h2.paragraph_format.space_after = Pt(6)
h2.paragraph_format.keep_with_next = True

h3 = doc.styles["Heading 3"]
style_font(h3, size=11.5, bold=True, color=RGBColor(0x33, 0x33, 0x33), cn=HEAD_CN, en=HEAD_CN)
h3.paragraph_format.space_before = Pt(10)
h3.paragraph_format.space_after = Pt(4)
h3.paragraph_format.keep_with_next = True

# 页脚页码
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_p)

doc.core_properties.title = "整理"
doc.core_properties.subject = "端侧 LLM 部署 · 量化 · 算子"
doc.core_properties.author = "renha"

# ---------------------------------------------------------------- 封面
for _ in range(3):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("整  理")
set_run_font(r, size=40, bold=True, color=C_DARK, cn=HEAD_CN, en=HEAD_CN)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
r = p.add_run("端侧 LLM · 量化 · 算子 · 推理系统")
set_run_font(r, size=13, color=C_MID, cn=HEAD_CN, en=HEAD_CN)

gap = doc.add_paragraph()
gap.paragraph_format.space_before = Pt(24)
bottom_border(gap, color="1F3864", size=6)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
r = p.add_run("“把模型拆到 op，把 op 拆到字节。”")
set_run_font(r, size=11, color=C_GRAY, italic=True)

for _ in range(6):
    doc.add_paragraph()

info = [
    ("使用者", "renha"),
    ("方向", "端侧推理 / 模型量化 / 算子优化"),
    ("版本", "v1.0"),
    ("起始", "2026-10"),
]
for k, v in info:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{k}　")
    set_run_font(r, size=10, color=C_GRAY)
    r = p.add_run(v)
    set_run_font(r, size=10, color=C_TEXT)

# ---------------------------------------------------------------- 使用说明
part(doc, "使用说明")
bottom_border(doc.paragraphs[-1])

blocks = [
    ("这本书怎么用", None),
    ("每个二级标题 = 一份 md 笔记。我在 md 里学、在这里归档；学到的东西先写成 md，讲顺了再粘进来，粘完把灰色提示行删掉。", "p"),
    ("只收「面试能讲、简历能写」的内容。全面性不是目标，能撑住 3 轮追问才是。", "p"),
    ("每个知识点必须带数字：形状、带宽、时延、精度掉点。没有数字的段落，面试官记不住。", "p"),
    ("每节末尾留一句「60 秒口述版」，用第一人称说，能对着镜子讲完再算学完。", "p"),
]
for text, kind in blocks:
    if kind is None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        set_run_font(r, size=11.5, bold=True, color=C_DARK, cn=HEAD_CN, en=HEAD_CN)
    else:
        doc.add_paragraph(text, style="List Bullet")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(14)
r = p.add_run("一句话人设")
set_run_font(r, size=11.5, bold=True, color=C_DARK, cn=HEAD_CN, en=HEAD_CN)

p = doc.add_paragraph()
p.paragraph_format.left_indent = Cm(0.6)
bottom_border(p, color="BFBFBF", size=4, space=6)
r = p.add_run(
    "我在高通做端侧 LLM 的部署与量化：把一个 Transformer 图从 PyTorch/ONNX 一路拆到 HTP 上的 op，"
    "在精度和时延之间做权衡，能定位到底是哪个 op、哪种量化策略让 W4A16 掉点，或者让时延卡在 memory-bound 上。"
)
set_run_font(r, size=10.5, color=RGBColor(0x40, 0x40, 0x40), italic=True)

# ---------------------------------------------------------------- 索引
doc.add_page_break()
doc.add_heading("索引与进度", level=1)
doc.paragraphs[-1].paragraph_format.space_before = Pt(0)
bottom_border(doc.paragraphs[-1])

rows = [
    ("01", "MobileNetV2 的 Graph Op 与 llama.cpp 落地", "二 · 算子", "notes/01_mobilenetv2_ggml.md", "已整理", ""),
    ("02", "Qwen3.5 在 transformers 里的实现", "三 · 架构", "notes/02_qwen35_transformer.md", "已整理", ""),
    ("03", "数值格式与硬件映射（int8/int4/fp8/mxfp4）", "一 · 量化", "", "待写", ""),
    ("04", "PTQ 全链路：校准、敏感层、混合精度", "一 · 量化", "", "待写", ""),
    ("05", "AWQ / GPTQ / SmoothQuant 的推导", "一 · 量化", "", "待写", ""),
    ("06", "旋转类量化：Hadamard / QuaRot / SpinQuant", "一 · 量化", "", "待写", ""),
    ("07", "QAT 与 STE：什么时候值得做", "一 · 量化", "", "待写", ""),
    ("08", "KV cache 量化与长上下文", "一 · 量化", "", "待写", ""),
    ("09", "Roofline：怎么判断一个 op 的瓶颈", "二 · 算子", "", "待写", ""),
    ("10", "int8 / W4A16 kernel 怎么写", "二 · 算子", "", "待写", ""),
    ("11", "PagedAttention 与 continuous batching", "三 · 架构", "", "待写", ""),
    ("12", "MoE / MLA / 稀疏注意力", "三 · 架构", "", "待写", ""),
    ("13", "端侧 VLM 与 VLA 的时延预算", "四 · 应用", "", "待写", ""),
    ("14", "RAG 的失败模式与评测", "四 · 应用", "", "待写", ""),
]
make_table(
    doc,
    ["#", "主题", "方向", "来源", "状态", "上次复习"],
    rows,
    widths=[1.0, 5.4, 2.2, 4.4, 1.6, 2.0],
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
r = p.add_run("状态用：待写 / 学习中 / 已整理 / 已讲顺（能 3 分钟口述）")
set_run_font(r, size=8.5, color=C_HINT, italic=True)

# ---------------------------------------------------------------- 第一部分
part(doc, "第一部分　模型量化")
bottom_border(doc.paragraphs[-1])

p = doc.add_paragraph()
r = p.add_run("主线：面试的主战场。目标不是「知道所有方法」，而是「能推公式 + 能选型 + 有数字」。")
set_run_font(r, size=10, color=C_GRAY, italic=True)

for name, note in [
    ("1.1　数值格式与硬件映射", "int8 / int4 / fp8 / mxfp4 / nvfp4；对称与非对称；per-tensor / per-channel / per-group；zero-point 在硬件上怎么省掉。"),
    ("1.2　PTQ 全链路", "校准集、MinMax / Percentile / MSE / KL、bias correction、逐层敏感度、混合精度搜索。"),
    ("1.3　高级 PTQ", "GPTQ（OBS + Hessian）、AWQ、SmoothQuant、AdaRound / BRECQ 的数学推导与实现差异。"),
    ("1.4　旋转类方法", "Hadamard / QuaRot / SpinQuant / PrefixQuant：为什么正交变换能压 outlier，代价在哪。"),
    ("1.5　QAT", "STE 的梯度偏差、LSQ / DoReFa、fake-quant 放置、蒸馏；QAT 相对 PTQ 的真实收益边界。"),
    ("1.6　评估与调优方法论", "ppl 与任务指标的背离、掉点归因、长尾与安全对齐、如何证明方案更好。"),
    ("1.7　部署侧特化", "W4A16 vs W8A8、unpack 开销、group size 与 kernel 对齐、KV 量化的额外读写。"),
]:
    section(doc, name)
    hint(doc, note)

# ---------------------------------------------------------------- 第二部分
part(doc, "第二部分　NPU / GPU 算子与 kernel 优化")
bottom_border(doc.paragraphs[-1])

p = doc.add_paragraph()
r = p.add_run("副线：从「部署工程师」往「算子/推理工程师」平移的通道。重在「能算清楚为什么快」。")
set_run_font(r, size=10, color=C_GRAY, italic=True)

for name, note in [
    ("2.1　计算图与 lowering", "ggml 的 tensor/op/graph/allocator；ONNX → 后端；op fusion；NCHW 与 NHWC 的代价。"),
    ("2.2　GEMM 与 tiling", "roofline、arithmetic intensity、shared memory / VTCM 分块、double buffering、寄存器分块。"),
    ("2.3　Memory-bound op", "DW conv、elementwise、softmax、RMSNorm：为什么慢，带宽上限怎么算。"),
    ("2.4　后端差异", "CPU(SIMD/多线程) / CUDA / HVX-HTP / Metal / Vulkan：同一个 op 的不同写法。"),
    ("2.5　量化 kernel", "int8 dp4a、W4A16 的 dequant-in-register、unpack 的指令开销。"),
    ("2.6　图级优化", "常量折叠、prefill 与 decode 的瓶颈切换、arena 内存复用、多后端调度与拷贝。"),
    ("2.7　Profiling", "nsys / ncu、llama-bench、ggml graph dump、HTP 侧 op 时间线。"),
    ("2.8　案例：MobileNetV2 的 op 图", "把 CNN 拆成 ggml 的 op 序列，算清每个 op 的算术强度，看 conv 被 lower 成 im2col + GEMM 的代价。"),
]:
    section(doc, name)
    if name.startswith("2.8"):
        hint(doc, "▸ 来源：notes/01_mobilenetv2_ggml.md（把 md 内容粘到此处后删除本行）")
    else:
        hint(doc, note)

# ---------------------------------------------------------------- 第三部分
part(doc, "第三部分　推理系统与模型架构")
bottom_border(doc.paragraphs[-1])

p = doc.add_paragraph()
r = p.add_run("补齐常识：不是主战场，但答不上来会显得不像做推理的。")
set_run_font(r, size=10, color=C_GRAY, italic=True)

for name, note in [
    ("3.1　推理引擎基础", "PagedAttention、continuous batching、prefix caching、chunked prefill、speculative decoding。"),
    ("3.2　调度与内存", "KV cache 管理 / swap / 抢占、TP / PP / EP、MoE 的 all-to-all。"),
    ("3.3　高效 attention 架构", "MLA、Gated DeltaNet / linear attention、稀疏注意力、混合层比例。"),
    ("3.4　KV 压缩与投机解码", "KV 量化、NativeKV / SPD 一类思路、EAGLE 的 draft 结构。"),
    ("3.5　端云协同", "端侧小模型 + 云端大模型、端侧 draft + 云端 verify、路由与降级。"),
    ("3.6　案例：Qwen3.5 的实现", "3:1 混合注意力、GDN 的递推、zero-centered RMSNorm、partial RoPE、状态与带宽的定量账。"),
]:
    section(doc, name)
    if name.startswith("3.6"):
        hint(doc, "▸ 来源：notes/02_qwen35_transformer.md（把 md 内容粘到此处后删除本行）")
    else:
        hint(doc, note)

# ---------------------------------------------------------------- 第四部分
part(doc, "第四部分　AI 应用")
bottom_border(doc.paragraphs[-1])

p = doc.add_paragraph()
r = p.add_run("差异化：决定你能不能把技术讲成产品故事，也是转机器人 / 具身的入口。不超过 10% 的精力。")
set_run_font(r, size=10, color=C_GRAY, italic=True)

for name, note in [
    ("4.1　RAG 与 Agentic RAG", "chunking、embedding、hybrid retrieval、rerank、query rewrite、tool use、评测指标。"),
    ("4.2　传统 CV 与端侧 CV", "检测 / 分割 / 跟踪、TFLite / QNN 部署、后处理开销。"),
    ("4.3　机器人与 VLA", "VLA 结构、action head、推理时延预算、sim2real。"),
    ("4.4　VLM 与 world model", "视觉 token 数控制、patch merge、动态分辨率、视频时序建模。"),
]:
    section(doc, name)
    hint(doc, note)

# ---------------------------------------------------------------- 附录
part(doc, "附录")
bottom_border(doc.paragraphs[-1])

doc.add_heading("A　面试问答总表", level=2)
hint(doc, "按「问题 / 一句话答案 / 追问点 / 我的项目证据」四列整理，面试前只复习这张表。")
make_table(
    doc,
    ["问题", "一句话答案", "可能的追问", "我的证据"],
    [("", "", "", ""), ("", "", "", ""), ("", "", "", "")],
    widths=[5.2, 5.6, 3.0, 2.8],
)

doc.add_heading("B　数字卡片", level=2)
hint(doc, "所有能背出来的数字放这里。面试时说出一个具体数字，可信度立刻不一样。")
make_table(
    doc,
    ["数字", "含义", "出处"],
    [
        ("3.47 M / ≈3 亿 MACs", "MobileNetV2 1.0@224 的参数量与计算量", "notes/01"),
        ("≈1 MAC/Byte", "MobileNetV2 depthwise 层的算术强度", "notes/01"),
        ("16384 元素/token", "Qwen3.5 每 token 的 KV（仅 8 个 full attention 层）", "notes/02"),
        ("134 MB", "Qwen3.5-9B 级配置在 4k 上下文的 KV（fp16）", "notes/02"),
        ("≈25 MB", "Qwen3.5 的 GDN 固定状态（24 层，fp16，与长度无关）", "notes/02"),
        ("≈768 token", "GDN 状态与 KV 带宽打平的上下文长度", "notes/02"),
    ],
    widths=[4.2, 8.4, 4.0],
)

doc.add_page_break()
doc.add_heading("C　项目复述模板", level=2)
hint(doc, "每个项目一段，严格四句：背景 → 我的动作 → 数字结果 → 学到什么。")
for text in [
    "项目名称 / 时间 / 团队规模",
    "背景：要解决什么问题，约束是什么（时延、内存、精度）",
    "我的动作：我具体做了什么，用了什么工具链，做了哪些取舍",
    "数字结果：before / after，至少一个百分比",
    "学到什么：一句反思，最好是「如果重做我会改什么」",
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("D　待办与复盘", level=2)
hint(doc, "每周写三行：本周学了什么、卡在哪、下周只做一件事。")
make_table(
    doc,
    ["周", "学了什么", "卡在哪", "下周只做一件事"],
    [("", "", "", ""), ("", "", "", ""), ("", "", "", ""), ("", "", "", ""), ("", "", "", "")],
    widths=[1.6, 5.4, 4.8, 4.8],
)

doc.save("整理.docx")
print("saved")
